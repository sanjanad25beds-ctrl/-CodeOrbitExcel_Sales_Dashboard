"""Tests for the Excel sales dashboard.

Verifies: KPI cards match independent recomputation from the sample CSV,
4 native charts exist on the Dashboard, the Data sheet is complete with
live Revenue/Profit formulas, and the Summary tables have the right shape.
"""
import pandas as pd
from openpyxl import load_workbook

XLSX = "Sales_Dashboard.xlsx"
CSV = "/home/hatch/workspace/your_files/codeorbit-internship/SalesPulse_AI/data/sample_sales.csv"


def _independent():
    df = pd.read_csv(CSV)
    df["Revenue"] = df["Quantity"] * df["Unit Price"] * (1 - df["Discount"])
    df["Profit"] = df["Revenue"] - df["Quantity"] * df["Unit Cost"]
    return df


def test_kpis_match_recomputation():
    df = _independent()
    wb = load_workbook(XLSX, data_only=False)
    db = wb["Dashboard"]
    vals = [db.cell(row=5, column=c).value for c in (1, 3, 5, 7, 9, 11)]
    assert abs(vals[0] - df["Revenue"].sum()) < 0.01
    assert abs(vals[1] - df["Profit"].sum()) < 0.01
    assert vals[2] == df["Order ID"].nunique()
    assert abs(vals[3] - df["Revenue"].sum() / df["Order ID"].nunique()) < 0.01
    assert abs(vals[4] - df["Profit"].sum() / df["Revenue"].sum()) < 1e-6
    assert vals[5] == df.groupby("Region")["Revenue"].sum().idxmax()


def test_four_charts_on_dashboard():
    wb = load_workbook(XLSX)
    charts = wb["Dashboard"]._charts
    assert len(charts) == 4
    kinds = sorted(type(c).__name__ for c in charts)
    assert kinds == ["BarChart", "BarChart", "LineChart", "PieChart"]


def test_data_sheet_complete_with_formulas():
    wb = load_workbook(XLSX)
    ws = wb["Data"]
    assert ws.max_row == 1508  # header + 1507 transactions
    assert ws.max_column == 15
    assert ws.cell(row=1, column=14).value == "Revenue"
    assert ws["N2"].value == "=H2*I2*(1-K2)"
    assert ws["O2"].value == "=N2-H2*J2"


def test_summary_tables_shape():
    wb = load_workbook(XLSX)
    sm = wb["Summary"]
    assert sm["A1"].value == "Monthly trend"
    # 12 monthly rows under the header (row 2) -> rows 3..14
    months = [sm.cell(row=r, column=1).value for r in range(3, 15)]
    assert len([m for m in months if m is not None]) == 12
